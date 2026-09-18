#!/usr/bin/env python3
"""Build the three-sheet test report workbook from verdicts.json.

  python build_report.py verdicts.json --out "Test Report.xlsx" --lang id

Sheets: Summary / Test Cases / Build Gates. Headers and verdict labels follow
--lang (id | en); the prose in the JSON is written out as-is, in whatever
language you wrote it.

The script REFUSES to write a workbook whose rows break the evidence rules:

  * PASS / FAIL / NA / SPEC_STALE need at least one evidence entry
    (file + line + quoted code),
  * NEEDS_SYSTEM needs concrete test steps,
  * SPEC_STALE needs notes explaining the delta.

Those are exactly the rows that get quietly waved through under time pressure,
which is why this is an error and not a warning. A row that genuinely cannot
carry evidence belongs under the BLOCKED verdict. --allow-incomplete downgrades
the check to warnings and lists them on the Summary sheet, for drafts only.
"""

import argparse
import datetime as _dt
import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl is required: pip install openpyxl\n")
    raise SystemExit(2)


def say(msg):
    sys.stdout.write(str(msg).encode("ascii", "replace").decode("ascii") + "\n")


VERDICTS = ["PASS", "FAIL", "NA", "NEEDS_SYSTEM", "SPEC_STALE", "BLOCKED"]

LABELS = {
    "id": {
        "PASS": "LULUS", "FAIL": "GAGAL", "NA": "N/A",
        "NEEDS_SYSTEM": "PERLU SISTEM", "SPEC_STALE": "SPEC BASI",
        "BLOCKED": "TERBLOKIR",
        "sheet_summary": "Ringkasan", "sheet_tests": "Unit Test",
        "sheet_gates": "Build Gates",
        "title": "Laporan Hasil Testing (Berbasis Source Code)",
        "tests_hdr": ["No", "Kasus Uji", "Cakupan", "Hasil Diharapkan", "Verdict",
                      "Bukti (file:baris)", "Kutipan Kode", "Catatan / Deviasi"],
        "gates_hdr": ["No", "Gerbang", "Verdict", "Bukti (file:baris)",
                      "Kutipan Kode", "Catatan"],
        "counts_hdr": ["Verdict", "Unit Test", "Build Gates", "Total"],
        "meta_hdr": ["Keterangan", "Nilai"],
        "sec_counts": "Rekap Verdict",
        "sec_stale": "Temuan SPEC BASI - spec tidak lagi cocok dengan kode",
        "sec_manual": "Antrean Uji Manual (PERLU SISTEM)",
        "sec_blocked": "Tidak Dapat Dinilai (TERBLOKIR) - source belum tersedia",
        "sec_issues": "Peringatan Kualitas Data",
        "col_item": "No", "col_what": "Item", "col_detail": "Penjelasan",
        "col_steps": "Langkah Uji di SAP",
        "meta_keys": {
            "run_date": "Tanggal run", "program": "Program", "tcode": "Transaction code",
            "spec_file": "File funcspec", "spec_sheet": "Sheet skenario",
            "spec_version": "Versi / tanggal spec", "source_root": "Folder source",
            "source_revision": "Revisi source yang diuji", "method": "Metode",
            "tester": "Penilai",
        },
        "method_default": "Pembacaan source code (tidak dijalankan di SAP)",
        "none": "(tidak ada)",
    },
    "en": {
        "PASS": "PASS", "FAIL": "FAIL", "NA": "N/A",
        "NEEDS_SYSTEM": "NEEDS SYSTEM", "SPEC_STALE": "SPEC STALE",
        "BLOCKED": "BLOCKED",
        "sheet_summary": "Summary", "sheet_tests": "Test Cases",
        "sheet_gates": "Build Gates",
        "title": "Test Result Report (Source-Code Based)",
        "tests_hdr": ["No", "Test Case", "Covers", "Expected Result", "Verdict",
                      "Evidence (file:line)", "Code Quote", "Notes / Deviation"],
        "gates_hdr": ["No", "Gate", "Verdict", "Evidence (file:line)",
                      "Code Quote", "Notes"],
        "counts_hdr": ["Verdict", "Test Cases", "Build Gates", "Total"],
        "meta_hdr": ["Item", "Value"],
        "sec_counts": "Verdict Tally",
        "sec_stale": "SPEC STALE findings - spec no longer matches the code",
        "sec_manual": "Manual Test Queue (NEEDS SYSTEM)",
        "sec_blocked": "Not Assessable (BLOCKED) - source not supplied",
        "sec_issues": "Data Quality Warnings",
        "col_item": "No", "col_what": "Item", "col_detail": "Explanation",
        "col_steps": "Test Steps in SAP",
        "meta_keys": {
            "run_date": "Run date", "program": "Program", "tcode": "Transaction code",
            "spec_file": "Funcspec file", "spec_sheet": "Scenario sheet",
            "spec_version": "Spec version / date", "source_root": "Source folder",
            "source_revision": "Source revision under test", "method": "Method",
            "tester": "Assessed by",
        },
        "method_default": "Source code reading (not executed in SAP)",
        "none": "(none)",
    },
}

FILLS = {
    "PASS": "C6EFCE", "FAIL": "FFC7CE", "NA": "E7E6E6",
    "NEEDS_SYSTEM": "FFEB9C", "SPEC_STALE": "E4D0F0", "BLOCKED": "D9D9D9",
}
FONTS = {
    "PASS": "006100", "FAIL": "9C0006", "NA": "595959",
    "NEEDS_SYSTEM": "9C6500", "SPEC_STALE": "5B2C6F", "BLOCKED": "3F3F3F",
}

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP_LEFT = Alignment(vertical="top", horizontal="left")
CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)


def pick(item, *names, default=""):
    for name in names:
        if item.get(name) not in (None, ""):
            return item[name]
    return default


def as_lines(value):
    if value in (None, ""):
        return []
    if isinstance(value, (list, tuple)):
        out = []
        for element in value:
            out.extend(as_lines(element))
        return out
    return [str(value).strip()]


def normalise(item, kind):
    verdict = str(pick(item, "verdict", default="")).strip().upper().replace(" ", "_")
    verdict = {"N/A": "NA", "SPEC-STALE": "SPEC_STALE",
               "NEEDS-SYSTEM": "NEEDS_SYSTEM"}.get(verdict, verdict)
    evidence = []
    for entry in item.get("evidence") or []:
        if isinstance(entry, str):
            evidence.append({"file": entry, "line": "", "code": ""})
            continue
        evidence.append({
            "file": str(pick(entry, "file", "path", default="")).strip(),
            "line": entry.get("line", ""),
            "code": str(pick(entry, "code", "snippet", "text", default="")).strip(),
        })
    return {
        "id": str(pick(item, "id", "no", default="")).strip(),
        "title": str(pick(item, "case", "gate", "title", "name", default="")).strip(),
        "covers": str(pick(item, "covers", "cakupan", "coverage", default="")).strip(),
        "expected": str(pick(item, "expected", "expected_result", default="")).strip(),
        "verdict": verdict,
        "evidence": evidence,
        "notes": str(pick(item, "notes", "note", "catatan", default="")).strip(),
        "steps": as_lines(pick(item, "test_steps", "steps", "langkah", default="")),
        "kind": kind,
    }


def validate(rows):
    problems = []
    seen = set()
    for row in rows:
        tag = row["id"] or "(missing id)"
        if not row["id"]:
            problems.append("a %s row has no id" % row["kind"])
        elif row["id"] in seen:
            problems.append("%s: duplicate id" % tag)
        seen.add(row["id"])

        if row["verdict"] not in VERDICTS:
            problems.append("%s: verdict %r is not one of %s"
                            % (tag, row["verdict"], "/".join(VERDICTS)))
            continue

        if row["verdict"] in ("PASS", "FAIL", "NA", "SPEC_STALE"):
            if not row["evidence"]:
                problems.append("%s: verdict %s with no evidence - cite file:line+code, "
                                "or use NEEDS_SYSTEM / BLOCKED"
                                % (tag, row["verdict"]))
            for ev in row["evidence"]:
                if not ev["file"] or ev["line"] in ("", None):
                    problems.append("%s: evidence entry missing file or line" % tag)
                elif not ev["code"]:
                    problems.append("%s: evidence %s:%s has no quoted code line"
                                    % (tag, ev["file"], ev["line"]))
        if row["verdict"] == "NEEDS_SYSTEM" and not row["steps"]:
            problems.append("%s: NEEDS_SYSTEM without test steps - write what a "
                            "tester should do and observe" % tag)
        if row["verdict"] == "SPEC_STALE" and not row["notes"]:
            problems.append("%s: SPEC_STALE without notes - explain spec vs code" % tag)
    return problems


def style_header(ws, row_no, count, start_col=1):
    for c in range(start_col, start_col + count):
        cell = ws.cell(row=row_no, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def paint_verdict(cell, verdict, label):
    cell.value = label
    cell.fill = PatternFill("solid", fgColor=FILLS.get(verdict, "FFFFFF"))
    cell.font = Font(bold=True, color=FONTS.get(verdict, "000000"), size=10)
    cell.alignment = CENTER
    cell.border = BORDER


def write_case_sheet(ws, rows, headers, L, with_covers):
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for row in rows:
        evidence = "\n".join(
            "%s:%s" % (e["file"], e["line"]) for e in row["evidence"]) or "-"
        quotes = "\n".join(e["code"] for e in row["evidence"] if e["code"]) or "-"
        notes = row["notes"]
        if row["steps"]:
            steps = "\n".join("%d. %s" % (i, s) for i, s in enumerate(row["steps"], 1))
            notes = (notes + "\n\n" if notes else "") + L["col_steps"] + ":\n" + steps
        if with_covers:
            values = [row["id"], row["title"], row["covers"], row["expected"],
                      "", evidence, quotes, notes]
        else:
            values = [row["id"], row["title"], "", evidence, quotes, notes]
        ws.append(values)
        r = ws.max_row
        verdict_col = 5 if with_covers else 3
        paint_verdict(ws.cell(row=r, column=verdict_col), row["verdict"],
                      L.get(row["verdict"], row["verdict"]))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if c != verdict_col:
                cell.alignment = TOP_WRAP
        ws.cell(row=r, column=1).alignment = TOP_LEFT

    widths = ([8, 46, 16, 46, 15, 26, 52, 52] if with_covers
              else [8, 60, 15, 26, 52, 58])
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)


def write_summary(ws, meta, tests, gates, L, lang, problems):
    def section(title):
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = SECTION_FONT
        return ws.max_row

    ws.append([L["title"]])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="1F3864")

    ws.append([])
    ws.append(L["meta_hdr"])
    style_header(ws, ws.max_row, 2)
    meta_keys = L["meta_keys"]
    ordered = ["run_date", "program", "tcode", "spec_file", "spec_sheet",
               "spec_version", "source_root", "source_revision", "method", "tester"]
    values = dict(meta)
    values.setdefault("run_date", _dt.date.today().isoformat())
    values.setdefault("method", L["method_default"])
    for key in ordered:
        if values.get(key):
            ws.append([meta_keys.get(key, key), str(values[key])])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=10)
            for c in (1, 2):
                ws.cell(row=ws.max_row, column=c).border = BORDER
                ws.cell(row=ws.max_row, column=c).alignment = TOP_WRAP
    for key, value in values.items():
        if key not in ordered and value:
            ws.append([meta_keys.get(key, key), str(value)])
            for c in (1, 2):
                ws.cell(row=ws.max_row, column=c).border = BORDER

    section(L["sec_counts"])
    ws.append(L["counts_hdr"])
    style_header(ws, ws.max_row, 4)
    tests_ref = "'%s'!$E$2:$E$%d" % (L["sheet_tests"], max(2, len(tests) + 1))
    gates_ref = "'%s'!$C$2:$C$%d" % (L["sheet_gates"], max(2, len(gates) + 1))
    for verdict in VERDICTS:
        label = L[verdict]
        ws.append([label,
                   '=COUNTIF(%s,"%s")' % (tests_ref, label),
                   '=COUNTIF(%s,"%s")' % (gates_ref, label),
                   None])
        r = ws.max_row
        ws.cell(row=r, column=4).value = "=B%d+C%d" % (r, r)
        paint_verdict(ws.cell(row=r, column=1), verdict, label)
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = CENTER
    ws.append([("TOTAL"), len(tests), len(gates), len(tests) + len(gates)])
    for c in range(1, 5):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True, size=10)
        ws.cell(row=ws.max_row, column=c).border = BORDER

    def listing(title, rows, detail_getter, detail_header):
        section(title)
        ws.append([L["col_item"], L["col_what"], detail_header])
        style_header(ws, ws.max_row, 3)
        if not rows:
            ws.append([L["none"], "", ""])
            for c in range(1, 4):
                ws.cell(row=ws.max_row, column=c).border = BORDER
            return
        for row in rows:
            ws.append([row["id"], row["title"], detail_getter(row)])
            for c in range(1, 4):
                ws.cell(row=ws.max_row, column=c).border = BORDER
                ws.cell(row=ws.max_row, column=c).alignment = TOP_WRAP

    everything = tests + gates
    listing(L["sec_stale"],
            [r for r in everything if r["verdict"] == "SPEC_STALE"],
            lambda r: r["notes"], L["col_detail"])
    listing(L["sec_manual"],
            [r for r in everything if r["verdict"] == "NEEDS_SYSTEM"],
            lambda r: "\n".join("%d. %s" % (i, s)
                                for i, s in enumerate(r["steps"], 1)),
            L["col_steps"])
    blocked = [r for r in everything if r["verdict"] == "BLOCKED"]
    if blocked:
        listing(L["sec_blocked"], blocked, lambda r: r["notes"], L["col_detail"])

    if problems:
        section(L["sec_issues"])
        for problem in problems:
            ws.append([problem])
            ws.cell(row=ws.max_row, column=1).alignment = TOP_WRAP

    for col, width in zip("ABCD", (34, 52, 82, 12)):
        ws.column_dimensions[col].width = width


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("verdicts_json")
    ap.add_argument("--out", required=True)
    ap.add_argument("--lang", choices=["id", "en"], default="en")
    ap.add_argument("--allow-incomplete", action="store_true",
                    help="downgrade evidence-rule errors to warnings (drafts only)")
    args = ap.parse_args()

    with open(args.verdicts_json, "r", encoding="utf-8") as fh:
        data = json.load(fh)

    L = LABELS[args.lang]
    meta = data.get("meta", {})
    tests = [normalise(x, "test") for x in data.get("unit_tests",
                                                    data.get("test_cases", []))]
    gates = [normalise(x, "gate") for x in data.get("build_gates", [])]

    problems = validate(tests + gates)
    if problems and not args.allow_incomplete:
        say("Refusing to write the workbook. %d row(s) break the evidence rules:"
            % len(problems))
        for problem in problems:
            say("  - " + problem)
        say("")
        say("Fix the rows, or rerun with --allow-incomplete for a draft (the")
        say("warnings are then listed on the Summary sheet).")
        raise SystemExit(1)

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = L["sheet_summary"]
    ws_tests = wb.create_sheet(L["sheet_tests"])
    ws_gates = wb.create_sheet(L["sheet_gates"])

    write_case_sheet(ws_tests, tests, L["tests_hdr"], L, with_covers=True)
    write_case_sheet(ws_gates, gates, L["gates_hdr"], L, with_covers=False)
    write_summary(ws_sum, meta, tests, gates, L, args.lang, problems)

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(args.out)

    say("wrote %s" % args.out)
    say("  %s: %d rows" % (L["sheet_tests"], len(tests)))
    say("  %s: %d rows" % (L["sheet_gates"], len(gates)))
    tally = {}
    for row in tests + gates:
        tally[row["verdict"]] = tally.get(row["verdict"], 0) + 1
    for verdict in VERDICTS:
        if tally.get(verdict):
            say("  %-13s %d" % (verdict, tally[verdict]))
    if problems:
        say("  %d data-quality warning(s) listed on the summary sheet" % len(problems))


if __name__ == "__main__":
    main()
